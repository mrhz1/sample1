"""Split the matched studies into preop / postop using REDCap surgery dates.

Answers the request for, per cohort:

    medical patients, PREOP only          3 modalities x 3 measures =  9
    surgical patients, PREOP and POSTOP   3 modalities x 3 measures = 18

where the three measures are, per modality (Echo, CT, MRI):

    a. matched raw scan and report       at least one study with its report
    b. raw scan, no matched report       has images, but none of them reported
    c. report, no raw scan               has a report, but no images at all

Read literally, the three are exclusive: a patient is counted once per
modality and phase, in the first of a / b / c that applies. A patient with one
reported echo and one unreported echo is (a) - they do have a matched pair. The
Studies sheet lists every row behind every number, so the other reading can be
recovered with a filter.

Surgery dates come from the REDCap export, which has one plain row per patient
and one extra '<x>_surgical_procedure' row per operation:

    en_sts_particid  redcap_repeat_instrument  ...  sxaa_sts_surgdt
    AA017
    AA017            aaoca_surgical_procedure  ...  8/12/2004 0:00
    AA017            aaoca_surgical_procedure  ...  8/13/2004 0:00

A patient with at least one surgery date is SURGICAL; one with none is
MEDICAL, and every study they have counts as preop. With several operations
the FIRST is the cut-off: preop means before any surgery. The column names
differ per cohort (sxaa_, sxav_), so they are found by suffix: '_particid' for
the code and '_surgdt' for the date.

Codes are matched on prefix + number, as report.py does, so REDCap's AA001 is
the archive's AA0001.

Reads report.py's study_report table; nothing is re-matched here.

Usage:
    python report.py --db inventory.db --prefixes AA,AVSD     # must run first
    python surgical_split.py --db inventory.db \\
        --surgery aaoca_redcap.xlsx avsd_redcap.xlsx --out surgical_split.xlsx

    --db PATH         Database, after report.py.
    --surgery FILES   One or more REDCap exports (.xlsx or .csv).
    --out PATH        Workbook to write (default surgical_split.xlsx).
    --surgery-day X   Where a study ON the day of the first surgery goes:
                      'pre' (default) or 'post'. Intra-operative TEEs land
                      here; the run prints how many rows it affects.
    --details         Also write the Studies sheet (every row behind every
                      number) and the Check sheet (what was not counted).
                      Off by default: the workbook is what gets sent.
"""

import argparse
import csv
import os
import re
import sqlite3
import sys
from collections import defaultdict
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="DDDDDD")

# status values written by report.py
S_BOTH = "report and image"
S_IMAGE = "only image"
S_REPORT = "only report"

# raw DICOM modality (or report.py's alias) -> the name the request uses
MODALITIES = [("US", "Echo"), ("CT", "CT"), ("MR", "MRI")]
MODALITY_NAME = dict(MODALITIES)

MEASURES = [
    ("a", "with a matched raw {m} and {m} report"),
    ("b", "with a raw {m}, but without a matched {m} report"),
    ("c", "without a raw {m}, but have {an} {m} report"),
]

PRE, POST = "preop", "postop"
UNDATED = "undated"
UNREADABLE = "surgery date unreadable"
NOT_IN_REDCAP = "not in REDCap"

CODE_RE = re.compile(r"^\s*([A-Za-z]+)0*(\d+)\s*$")


def code_key(text):
    """'AA001' and 'AA0001' -> ('AA', 1). None if it is not a code."""
    m = CODE_RE.match(str(text or ""))
    return (m.group(1).upper(), int(m.group(2))) if m else None


def parse_date(value):
    """A REDCap date cell -> datetime.date, or None if empty / unreadable.

    Excel hands back a datetime; a CSV export gives '2009-12-30 00:00' (raw)
    or '12/30/2009 0:00' (as displayed, US order).
    """
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().split(" ")[0]
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def read_rows(path):
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as fh:
            return list(csv.reader(fh))
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    wb.close()
    return rows


def find_column(header, suffix, path):
    hits = [i for i, h in enumerate(header)
            if str(h or "").strip().lower().endswith(suffix)]
    if len(hits) != 1:
        sys.exit(f"{path}: expected one column ending '{suffix}', found "
                 f"{len(hits)} in {header}")
    return hits[0]


def load_surgery(paths):
    """code -> sorted surgery dates (empty list = medical), the raw id as
    REDCap spells it, every value that could not be read, and the codes those
    belong to - which are left out of the counts rather than guessed at."""
    surgeries, spelling, problems, unreadable = {}, {}, [], set()
    for path in paths:
        rows = read_rows(path)
        if not rows:
            sys.exit(f"{path}: empty")
        header = rows[0]
        id_col = find_column(header, "_particid", path)
        date_col = find_column(header, "_surgdt", path)
        for n, row in enumerate(rows[1:], start=2):
            raw = row[id_col] if id_col < len(row) else None
            if raw in (None, ""):
                continue
            key = code_key(raw)
            if key is None:
                problems.append((os.path.basename(path), n, raw, "", "not a patient code"))
                continue
            spelling.setdefault(key, str(raw).strip())
            dates = surgeries.setdefault(key, [])
            cell = row[date_col] if date_col < len(row) else None
            when = parse_date(cell)
            if when:
                dates.append(when)
            elif cell not in (None, ""):
                unreadable.add(key)
                problems.append((os.path.basename(path), n, raw, str(cell),
                                 "surgery date not readable"))
    for dates in surgeries.values():
        dates.sort()
    return surgeries, spelling, problems, unreadable


def study_day(text):
    """A study_report date -> datetime.date, or None.

    report.py writes YYYYMMDD, but databases built by other versions hold
    2016-01-25, so only the digits are trusted.
    """
    digits = re.sub(r"\D", "", str(text or ""))[:8]
    try:
        return datetime.strptime(digits, "%Y%m%d").date()
    except ValueError:
        return None


def phase_of(study_date, first_surgery, surgery_day):
    when = study_day(study_date)
    if when is None:
        return UNDATED
    if first_surgery is None:
        return PRE
    if when < first_surgery:
        return PRE
    if when > first_surgery:
        return POST
    return PRE if surgery_day == "pre" else POST


def measure_of(statuses):
    if S_BOTH in statuses:
        return "a"
    if S_IMAGE in statuses:
        return "b"
    return "c"


def sheet(wb, title, header, widths, data, freeze=True):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in data:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    if freeze:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{ws.max_row}"
    return ws


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--db", default="inventory.db")
    ap.add_argument("--surgery", nargs="+", required=True)
    ap.add_argument("--out", default="surgical_split.xlsx")
    ap.add_argument("--surgery-day", choices=("pre", "post"), default="pre")
    ap.add_argument("--details", action="store_true")
    args = ap.parse_args()

    if not os.path.exists(args.db):
        sys.exit(f"no such database: {args.db}")
    conn = sqlite3.connect(f"file:{os.path.abspath(args.db)}?mode=ro", uri=True)
    if not conn.execute("SELECT 1 FROM sqlite_master WHERE name = 'study_report'"
                        ).fetchone():
        sys.exit("no study_report table - run report.py --prefixes ... first")

    surgeries, spelling, problems, unreadable = load_surgery(args.surgery)
    cohorts = sorted({p for p, _ in surgeries})

    # Every study / report row for a patient in one of the REDCap cohorts.
    rows, archive_codes, on_the_day = [], {}, 0
    for code, sdate, modality, slices, report, status, note, folder in conn.execute(
            "SELECT code, study_date, modality, slices, report, status, note,"
            " folder FROM study_report WHERE code IS NOT NULL"
            " AND status IN (?, ?, ?)", (S_BOTH, S_IMAGE, S_REPORT)):
        key = code_key(code)
        if key is None or key[0] not in cohorts:
            continue
        archive_codes[key] = code
        if key not in surgeries:
            phase, first = NOT_IN_REDCAP, None
        elif key in unreadable:
            # A bad date must not quietly turn a surgical patient medical.
            phase, first = UNREADABLE, None
        else:
            first = surgeries[key][0] if surgeries[key] else None
            phase = phase_of(sdate, first, args.surgery_day)
            if first and study_day(sdate) == first:
                on_the_day += 1
        rows.append(dict(key=key, code=code, date=sdate or "",
                         modality=modality or "", slices=slices or 0,
                         report=report or "", status=status, note=note or "",
                         folder=folder or "", phase=phase, first=first))
    conn.close()

    # (key, modality, phase) -> statuses seen; one patient, one verdict.
    seen = defaultdict(set)
    for r in rows:
        if r["modality"] in MODALITY_NAME and r["phase"] in (PRE, POST):
            seen[(r["key"], r["modality"], r["phase"])].add(r["status"])
    verdict = {k: measure_of(v) for k, v in seen.items()}

    def group(key):
        if key in unreadable:
            return "check"
        return "surgical" if surgeries.get(key) else "medical"

    # cohort -> (group, phase, modality, measure) -> patient count
    counts = defaultdict(lambda: defaultdict(int))
    for (key, modality, phase), m in verdict.items():
        counts[key[0]][(group(key), phase, modality, m)] += 1

    wb = Workbook()
    wb.remove(wb.active)

    # --- Summary: laid out the way the request asks for it ---------------
    summary = []
    for cohort in cohorts:
        summary.append([f"{cohort} cohort", "", "", ""])
        for raw, name in MODALITIES:
            for letter, text in MEASURES:
                c = counts[cohort]
                summary.append([
                    f"  {letter}. {text.format(m=name, an='an' if name[0] in 'AEIMO' else 'a')}",
                    c[("medical", PRE, raw, letter)],
                    c[("surgical", PRE, raw, letter)],
                    c[("surgical", POST, raw, letter)]])
        summary.append(["", "", "", ""])
    ws = sheet(wb, "Summary",
               ["Number of patients", "Medical - PREOP", "Surgical - PREOP",
                "Surgical - POSTOP"], (52, 16, 16, 16), summary, freeze=False)
    for row in ws.iter_rows(min_row=2):
        if row[0].value and row[0].value.endswith("cohort"):
            row[0].font = Font(bold=True)

    # --- Patients: one line each, the letter behind every number ---------
    columns = [(raw, ph) for raw, _ in MODALITIES for ph in (PRE, POST)]
    pdata = []
    for key in sorted(surgeries):
        dates = surgeries[key]
        line = [archive_codes.get(key, ""), spelling[key], group(key),
                len(dates), dates[0] if dates else None,
                ", ".join(d.isoformat() for d in dates)]
        for raw, ph in columns:
            line.append(verdict.get((key, raw, ph), ""))
        pdata.append(line)
    ws = sheet(wb, "Patients",
          ["Archive Code", "REDCap ID", "Group", "Surgeries",
           "First Surgery", "All Surgery Dates"]
          + [f"{MODALITY_NAME[raw]} {ph}" for raw, ph in columns],
          (13, 12, 10, 10, 13, 30) + (11,) * len(columns), pdata)
    for cell in ws["E"][1:]:
        cell.number_format = "yyyy-mm-dd"

    # --- Check: what could not be placed ---------------------------------
    check = [list(p) for p in problems]
    for key in sorted(set(archive_codes) - set(surgeries)):
        check.append(["archive", "", archive_codes[key], "",
                      "in the archive but not in REDCap - not counted"])
    undated = defaultdict(int)
    for r in rows:
        if r["phase"] == UNDATED:
            undated[r["code"]] += 1
    for code, n in sorted(undated.items()):
        check.append(["archive", "", code, "",
                      f"{n} row(s) with no date - not counted"])
    # The two sheets below are for checking the numbers, not for sending,
    # so they are only written on request.
    if args.details:
        # --- Studies: every row behind every number --------------------------
        sdata = []
        for r in sorted(rows, key=lambda r: (r["key"], r["date"])):
            days = ""
            if r["first"] and study_day(r["date"]):
                days = (study_day(r["date"]) - r["first"]).days
            sdata.append([r["code"], group(r["key"]) if r["key"] in surgeries else "",
                          r["first"], r["date"], days,
                          MODALITY_NAME.get(r["modality"], r["modality"]),
                          r["phase"], r["status"], r["slices"], r["report"],
                          r["note"], r["folder"]])
        ws = sheet(wb, "Studies",
              ["Patient Code", "Group", "First Surgery", "Study Date",
               "Days From Surgery", "Modality", "Phase", "Status", "DICOM Files",
               "Report", "Note", "Folder"],
              (13, 10, 13, 11, 11, 10, 13, 17, 11, 36, 40, 60), sdata)
        for cell in ws["C"][1:]:
            cell.number_format = "yyyy-mm-dd"


        sheet(wb, "Check", ["Source", "Row", "Code", "Value", "Problem"],
              (22, 6, 13, 20, 50), check)


    wb.save(args.out)

    print(f"wrote {args.out}")
    for cohort in cohorts:
        members = [k for k in surgeries if k[0] == cohort]
        n = defaultdict(int)
        for k in members:
            n[group(k)] += 1
        print(f"  {cohort}: {n['medical']} medical, {n['surgical']} surgical"
              + (f", {n['check']} with an unreadable surgery date" if n["check"] else ""))
    print(f"  {on_the_day:,} rows on the day of first surgery, counted as "
          f"{'preop' if args.surgery_day == 'pre' else 'postop'}"
          f" (--surgery-day)")
    if check:
        print(f"  {len(check):,} items not counted - see the Check sheet"
              + ("" if args.details else " (add --details)"))


if __name__ == "__main__":
    main()
