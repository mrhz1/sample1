"""A two-sheet workbook: the headline numbers, and one line per patient.

Answers "what exactly does AA0001 have?" in a single row rather than across
several sheets:

    AA0001   600 DICOM images, 7 PDFs
             400 DICOMs have a report, 200 do not
             3 reports have no DICOM | 610 files, 1.2 GB

The counts here are **DICOM files**, not studies. master_report.xlsx counts
studies, because that is the unit a report is written about; this counts the
files underneath them, because that is the unit the archive is measured in. A
patient with one reported study of 400 slices and one unreported study of 200
appears here as 400 and 200, and there as 1 and 1. Both are right; this is the
one people mean when they ask how much is outstanding.

Reads the matching report.py already wrote to the database, so the two files
can never disagree and nothing is re-matched here.

Usage:
    python report.py --db inventory.db --prefixes AA   # must run first
    python patient_summary.py --db inventory.db --out patient_summary.xlsx

    --db PATH        Database, after crawl.py, probe.py and report.py.
    --out PATH       Workbook to write (default patient_summary.xlsx).
    --prefix LIST    Only these prefixes, comma-separated.
"""

import argparse
import os
import re
import sqlite3
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="DDDDDD")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")

# status values written by report.py
S_BOTH = "report and image"
S_IMAGE = "only image"
S_REPORT = "only report"
S_OTHER = "other pdf"

# Spelled out, no abbreviations - these headings get read by people who do not
# work with the archive daily, and "DICOMs w/o Report" means nothing to them.
HEADER = [
    ("Patient Code", 14),
    ("Total DICOM Images", 17),
    ("Total PDFs", 12),
    ("DICOMs With a Report", 18),
    ("DICOMs With No Report", 18),
    ("Reports With No DICOM", 18),
    ("Other PDFs", 12),
    ("Total Files (All Types)", 18),
    ("Total Size", 13),
]


def human_bytes(n):
    size = float(n or 0)
    for unit in ("B", "KB", "MB", "GB", "TB", "PB"):
        if size < 1024 or unit == "PB":
            return f"{size:,.1f} {unit}"
        size /= 1024


def code_prefix(code):
    m = re.match(r"[A-Za-z]+", code or "")
    return m.group(0).upper() if m else ""


def collect(conn):
    """code -> every number on the patient line."""
    rows = {}

    def row(code):
        return rows.setdefault(code, {
            "dicom_reported": 0, "reports_covering": set(),
            "dicom_unreported": 0, "reports_orphan": 0, "other_pdf": 0,
            "other_files": 0, "dicom_files": 0, "pdf_files": 0,
            "total_files": 0, "size": 0})

    for code, status, slices, report in conn.execute(
            "SELECT code, status, slices, report FROM study_report"
            " WHERE code IS NOT NULL"):
        r = row(code)
        if status == S_BOTH:
            r["dicom_reported"] += slices or 0
            if report:
                r["reports_covering"].add(report)
        elif status == S_IMAGE:
            r["dicom_unreported"] += slices or 0
        elif status == S_REPORT:
            r["reports_orphan"] += 1
        elif status == S_OTHER:
            r["other_pdf"] += 1

    for code, kind, n, size in conn.execute(
            "SELECT code, kind, COUNT(*), COALESCE(SUM(size), 0) FROM files"
            " WHERE code IS NOT NULL GROUP BY code, kind"):
        r = row(code)
        r["total_files"] += n
        r["size"] += size
        if kind == "dicom":
            r["dicom_files"] += n
        elif kind == "pdf":
            r["pdf_files"] += n
        else:
            r["other_files"] += n
    return rows


def build_rows(rows):
    """Totals first, then the split - so a reader can see the whole before
    the parts, and spot it when the parts do not add up to it."""
    out = []
    for code in sorted(rows):
        r = rows[code]
        out.append([code, r["dicom_files"], r["pdf_files"],
                    r["dicom_reported"], r["dicom_unreported"],
                    r["reports_orphan"], r["other_pdf"],
                    r["total_files"], human_bytes(r["size"])])
    return out


def overview(conn, rows, root):
    total = lambda k: sum(r[k] for r in rows.values())  # noqa: E731
    reported = total("dicom_reported")
    unreported = total("dicom_unreported")
    both = sum(1 for r in rows.values()
               if (r["dicom_reported"] or r["dicom_unreported"])
               and (r["reports_covering"] or r["reports_orphan"]))
    only_img = sum(1 for r in rows.values()
                   if (r["dicom_reported"] or r["dicom_unreported"])
                   and not (r["reports_covering"] or r["reports_orphan"]))
    only_rep = sum(1 for r in rows.values()
                   if not (r["dicom_reported"] or r["dicom_unreported"])
                   and (r["reports_covering"] or r["reports_orphan"]))
    neither = len(rows) - both - only_img - only_rep
    pct = f"{100.0 * reported / (reported + unreported):.1f}%" if (
        reported + unreported) else "-"
    return [
        ("Archive root", root),
        ("Patients", len(rows)),
        ("", ""),
        ("Image files (DICOM) that have a report", reported),
        ("Image files with no report", unreported),
        ("Share of image files reported", pct),
        ("Image files not in any study", sum(
            max(r["dicom_files"] - r["dicom_reported"] - r["dicom_unreported"], 0)
            for r in rows.values())),
        ("", ""),
        ("Reports that cover images", sum(
            len(r["reports_covering"]) for r in rows.values())),
        ("Reports with no images", total("reports_orphan")),
        ("Other PDFs (no date - probably not reports)", total("other_pdf")),
        ("Other files (not images, not PDFs)", total("other_files")),
        ("", ""),
        ("Patients with both images and reports", both),
        ("Patients with images but no reports", only_img),
        ("Patients with reports but no images", only_rep),
        ("Patients with neither images nor reports", neither),
        ("", ""),
        ("Total files", total("total_files")),
        ("Total size", human_bytes(total("size"))),
    ]


def sheet(wb, title, header, widths, data, freeze=True):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in data:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    if freeze:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{ws.max_row}"
    return ws


def attribution_banner(conn):
    """When the codes in this database were written, and by what command.

    Printed because these tools display report.py's attribution rather than
    deriving their own - so a stale database looks exactly like a bug in the
    tool that reads it.
    """
    meta = dict(conn.execute(
        "SELECT key, value FROM meta WHERE key IN"
        " ('report_run_at', 'report_args')"))
    when = meta.get("report_run_at")
    if not when:
        return ("codes were written by a report.py run older than this stamp -"
                " re-run it if anything below looks stale")
    return f"codes written by report.py at {when}  ({meta.get('report_args', '')})"


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--db", default="inventory.db")
    ap.add_argument("--out", default="patient_summary.xlsx")
    ap.add_argument("--prefix", default=None)
    args = ap.parse_args()

    if not os.path.exists(args.db):
        sys.exit(f"no such database: {args.db}")
    conn = sqlite3.connect(f"file:{os.path.abspath(args.db)}?mode=ro", uri=True)
    have = {r[0] for r in conn.execute(
        "SELECT name FROM sqlite_master WHERE type = 'table'")}
    if "study_report" not in have:
        sys.exit("no study_report table - run report.py --prefixes ... first")

    root = conn.execute("SELECT value FROM meta WHERE key='root'").fetchone()[0]
    banner = attribution_banner(conn)
    rows = collect(conn)
    if args.prefix:
        wanted = {p.strip().upper() for p in args.prefix.split(",") if p.strip()}
        rows = {c: r for c, r in rows.items() if code_prefix(c) in wanted}
        if not rows:
            print(f"warning: no codes with prefix {args.prefix}", file=sys.stderr)
    conn.close()

    wb = Workbook()
    wb.remove(wb.active)
    ws = sheet(wb, "Overview", ["Measure", "Value"], (44, 60),
               overview(None, rows, root), freeze=False)
    ws.auto_filter.ref = None

    data = build_rows(rows)
    ws = sheet(wb, "Patients", [h for h, _ in HEADER], [w for _, w in HEADER],
               data)
    # Anything outstanding on a row is worth the eye landing on it.
    for row in ws.iter_rows(min_row=2):
        if (row[4].value or 0) or (row[5].value or 0):
            for cell in row:
                cell.fill = WARN_FILL
    wb.save(args.out)

    print(f"wrote {args.out}")
    print(f"  {banner}")
    print(f"  {len(data):,} patients")
    print(f"  {sum(r[3] for r in data):,} DICOMs have a report, "
          f"{sum(r[4] for r in data):,} do not")
    print(f"  {sum(r[5] for r in data):,} reports with no DICOM, "
          f"{sum(r[6] for r in data):,} other PDFs")


if __name__ == "__main__":
    main()
