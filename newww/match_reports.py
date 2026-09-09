"""
Match DICOM image studies (images_root/<code folder>/...) against PDF reports
(reports_root/<code folder>/...) by study date + modality, and write one Excel
results file per code plus a combined summary.

Layout expected:
    <images_root>/AVSD0001 base date/...   <- DICOM files, any extension/none
    <reports_root>/AVSD0001 follow up/...  <- PDF reports

Folder names on either side may carry extra text around the code
("AVSD0001", "AVSD0001 base date", "AVSD_0001 follow-up"); only the code
itself is used to pair an images folder with a reports folder, and several
folders sharing one code are scanned together as that code.

PDF file names no longer have to contain the code - the code always comes
from the folder the PDF sits in. Dates in PDF names may be numeric
(22-11-2018) or use a month name (22-NOV-2018), and the modality is read
from words like Echo / MRI / CT in the PDF name (falling back to the folder
name when the file name has neither a date nor a modality).

Usage:
    python match_reports.py <images_root> <reports_root> [output.xlsx] [options]

Positional args:
    images_root    Folder containing one subfolder per code with DICOM files
    reports_root   Folder containing one subfolder per code with PDF reports
    output.xlsx    Combined summary path (default: ./match_report.xlsx)

Options:
    --workers N       Run N codes in parallel (one worker process per code).
                      Default 1 (serial, original behavior). Safe to raise since
                      every code reads/writes independently of the others.
    --code LIST       Only process these codes, comma-separated
                      (e.g. --code AVSD0001,AVSD0002). Default: all codes found.
                      --rcode is accepted as an alias.
    --code-regex RE   Pattern that finds the code inside a folder name.
                      Default: AVSD[-_ ]?\\d+ (case-insensitive).
    --date-order X    Order to assume for ambiguous all-numeric dates such as
                      05-06-2018: dmy (default, = 5 June) or mdy (= 6 May).
                      Dates where one number is > 12, ISO dates and month-name
                      dates are detected regardless of this setting.
    --force           Reprocess codes even if a <code>-results.xlsx already
                      exists (default: skip codes already done, for resuming).
    --debug           Verbose per-file parse/skip logging to stdout.

Examples:
    # Basic run, serial
    python match_reports.py "/mnt/data/images" "/mnt/data/reports" match_report.xlsx

    # Parallel across codes, 8 workers
    python match_reports.py "/mnt/data/images" "/mnt/data/reports" match_report.xlsx --workers 8

    # Resume an interrupted run (already-done codes are skipped automatically)
    python match_reports.py "/mnt/data/images" "/mnt/data/reports" match_report.xlsx --workers 8

    # Reprocess only specific codes, forcing even if already done
    python match_reports.py "/mnt/data/images" "/mnt/data/reports" match_report.xlsx --code AVSD0001,AVSD0002 --force

Per-code results land in <output.xlsx's folder>/results/<code>-results.xlsx;
progress/heartbeat logging is appended to results/progress.log on every run.
"""

import concurrent.futures
import os
import re
import sys
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import pydicom
from openpyxl import Workbook, load_workbook

# --- Study code -------------------------------------------------------------
# Folder names look like "AVSD0001", "AVSD0001 base date", "AVSD_0001 follow up".
# Only the code part is used to pair an images folder with a reports folder.
DEFAULT_CODE_PATTERN = r"AVSD[-_ ]?\d+"

# --- Dates ------------------------------------------------------------------
MONTH_NUMBERS = {
    "JAN": "01", "FEB": "02", "MAR": "03", "APR": "04",
    "MAY": "05", "JUN": "06", "JUL": "07", "AUG": "08",
    "SEP": "09", "OCT": "10", "NOV": "11", "DEC": "12",
}
_MONTH_ALT = "|".join(MONTH_NUMBERS)
_SEP = r"[-_.,\s]*"

# 2018-11-22 / 2018_11_22
ISO_DATE_RE = re.compile(r"(?<!\d)(\d{4})[-_.](\d{1,2})[-_.](\d{1,2})(?!\d)")
# 22-NOV-2018, 22 November 2018, 22NOV2018
DAY_MONTH_NAME_RE = re.compile(
    rf"(?<!\d)(\d{{1,2}}){_SEP}({_MONTH_ALT})[A-Z]*{_SEP}(\d{{4}})(?!\d)"
)
# NOV-22-2018, November 22, 2018
MONTH_NAME_DAY_RE = re.compile(
    rf"(?<![A-Z])({_MONTH_ALT})[A-Z]*{_SEP}(\d{{1,2}}){_SEP}(\d{{4}})(?!\d)"
)
# 22-11-2018 / 11-22-2018 / 22.11.2018 - order resolved by resolve_numeric_date
NUMERIC_DATE_RE = re.compile(r"(?<!\d)(\d{1,2})[-_.](\d{1,2})[-_.](\d{4})(?!\d)")

# How often, in files examined, to emit a heartbeat line while scanning a
# single code folder. Large folders otherwise look hung for hours.
PROGRESS_EVERY = 2000

SUMMARY_HEADER = [
    "Code",
    "Study Date",
    "Modality",
    "DICOM File Count",
    "Matched PDF File(s)",
    "Match Status",
]
SUMMARY_WIDTHS = (12, 14, 10, 16, 45, 40)

# Characters that can't go in a filename, in case a code folder has one.
UNSAFE_NAME_RE = re.compile(r'[\\/:*?"<>|]')

# PDF-filename token (uppercased) -> standard DICOM Modality code.
# Extend this as you find more variants in the real filenames.
MODALITY_ALIASES = {
    "CT": "CT",
    "CTA": "CT",
    "CT ANGIO": "CT",
    "MRI": "MR",
    "MR": "MR",
    "MRA": "MR",
    "CMR": "MR",
    "CARDIAC MRI": "MR",
    "ECHO": "US",
    "ECHOCARDIOGRAM": "US",
    "ECHOCARDIOGRAPHY": "US",
    "TTE": "US",
    "TEE": "US",
    "ULTRASOUND": "US",
    "US": "US",
    "XRAY": "CR",
    "X-RAY": "CR",
    "CXR": "CR",
    "DX": "DX",
    "MAMMO": "MG",
    "MG": "MG",
    "PET": "PT",
    "SPECT": "NM",
    "NM": "NM",
    "ANGIO": "XA",
    "ANGIOGRAM": "XA",
    "ANGIOGRAPHY": "XA",
    "CATH": "XA",
    "ECG": "ECG",
    "EKG": "ECG",
}
# Longest alias first so "MRI" wins over "MR" and "ECHOCARDIOGRAM" over "ECHO".
MODALITY_RE = re.compile(
    r"(?<![A-Z0-9])("
    + "|".join(
        re.escape(k) for k in sorted(MODALITY_ALIASES, key=len, reverse=True)
    )
    + r")(?![A-Z0-9])"
)

# Raw DICOM Modality tag value (uppercased) -> human-readable label, for
# display in the Excel outputs only. Real-world files don't always use the
# standard 2-letter code (e.g. some write "Ultrasonic" instead of "US") -
# extend this as you encounter more raw values in the real data.
MODALITY_DISPLAY_NAMES = {
    "CT": "CT",
    "MR": "MRI",
    "US": "Echo",
    "ULTRASOUND": "Echo",
    "ULTRASONIC": "Echo",
    "CR": "X-Ray",
    "DX": "X-Ray",
    "MG": "Mammography",
    "PT": "PET",
    "NM": "Nuclear Medicine",
    "XA": "Angiography",
    "RF": "Fluoroscopy",
    "SR": "Structured Report",
    "ECG": "ECG",
}

_log_fh = None


def open_log(path):
    """Append-mode log so successive (resumed) runs stack up in one place."""
    global _log_fh
    _log_fh = open(path, "a", encoding="utf-8")
    _log_fh.write(f"\n===== run started {datetime.now():%Y-%m-%d %H:%M:%S} =====\n")
    _log_fh.flush()


def _init_worker(log_path):
    """ProcessPoolExecutor initializer: give each worker process its own
    handle to the same append-mode log file (small writes are atomic under
    O_APPEND, so interleaved lines from multiple workers don't corrupt)."""
    global _log_fh
    _log_fh = open(log_path, "a", encoding="utf-8")


def log(message):
    line = f"[{datetime.now():%Y-%m-%d %H:%M:%S}] {message}"
    print(line, flush=True)
    if _log_fh is not None:
        _log_fh.write(line + "\n")
        _log_fh.flush()


def format_duration(seconds):
    seconds = int(seconds)
    hours, rest = divmod(seconds, 3600)
    minutes, secs = divmod(rest, 60)
    if hours:
        return f"{hours}h{minutes:02d}m{secs:02d}s"
    if minutes:
        return f"{minutes}m{secs:02d}s"
    return f"{secs}s"


def display_modality(code):
    if not code:
        return code
    return MODALITY_DISPLAY_NAMES.get(code.strip().upper(), code)


def result_path(results_dir, code):
    return results_dir / f"{UNSAFE_NAME_RE.sub('_', code)}-results.xlsx"


def extract_code(name, code_re):
    """Pull the study code out of a folder name, ignoring any extra text.

    "AVSD0001 base date" / "avsd_0001 follow-up" -> "AVSD0001".
    Returns None when the name carries no code at all.
    """
    match = code_re.search(name)
    if not match:
        return None
    raw = match.group(0)
    digits = re.search(r"\d+", raw).group(0)
    prefix = raw[: raw.index(digits)].strip(" -_").upper()
    return f"{prefix}{digits}"


def _make_date(year, month, day):
    """Validate a Y/M/D triple and render it the way DICOM StudyDate does."""
    month, day = int(month), int(day)
    if not (1 <= month <= 12 and 1 <= day <= 31):
        return None
    return f"{int(year):04d}{month:02d}{day:02d}"


def resolve_numeric_date(first, second, year, date_order):
    """22-11-2018 -> 20181122. Whichever number is > 12 has to be the day;
    when both could be a month, fall back to the configured order."""
    a, b = int(first), int(second)
    if a > 12 and b <= 12:
        return _make_date(year, b, a)
    if b > 12 and a <= 12:
        return _make_date(year, a, b)
    if date_order == "mdy":
        return _make_date(year, a, b)
    return _make_date(year, b, a)


def parse_pdf_date(name, date_order="dmy"):
    """First date found in a file/folder name, as YYYYMMDD (or None)."""
    upper = name.upper()

    match = ISO_DATE_RE.search(upper)
    if match:
        year, month, day = match.groups()
        date = _make_date(year, month, day)
        if date:
            return date

    match = DAY_MONTH_NAME_RE.search(upper)
    if match:
        day, month_name, year = match.groups()
        date = _make_date(year, MONTH_NUMBERS[month_name], day)
        if date:
            return date

    match = MONTH_NAME_DAY_RE.search(upper)
    if match:
        month_name, day, year = match.groups()
        date = _make_date(year, MONTH_NUMBERS[month_name], day)
        if date:
            return date

    match = NUMERIC_DATE_RE.search(upper)
    if match:
        first, second, year = match.groups()
        return resolve_numeric_date(first, second, year, date_order)

    return None


def parse_pdf_modality(name):
    match = MODALITY_RE.search(name.upper())
    if not match:
        return None
    return MODALITY_ALIASES[match.group(1)]


def name_chain(path, root):
    """The PDF's own name, then every folder name up to and including the
    code folder - used to recover a date/modality the file name lacks."""
    names = [path.name]
    parent = path.parent
    while True:
        names.append(parent.name)
        if parent == root or parent.parent == parent:
            break
        parent = parent.parent
    return names


def read_dicom_study(path, debug=False):
    try:
        ds = pydicom.dcmread(str(path), stop_before_pixels=True, force=False)
    except Exception as exc:
        if debug:
            forced_ok = False
            forced_date = None
            try:
                ds2 = pydicom.dcmread(str(path), stop_before_pixels=True, force=True)
                forced_ok = True
                forced_date = str(ds2.get("StudyDate", ""))
            except Exception:
                pass
            hint = ""
            if forced_ok:
                hint = (
                    f" -> READABLE with force=True (StudyDate={forced_date!r}); "
                    "this file is likely missing the DICOM preamble/DICM header"
                )
            print(f"  [SKIP] {path}: {type(exc).__name__}: {exc}{hint}")
        return None
    study_date, modality = str(ds.get("StudyDate", "")), str(ds.get("Modality", ""))
    if debug:
        print(f"  [OK]   {path}: StudyDate={study_date!r} Modality={modality!r}")
    return study_date, modality


def scan_reports_dir(code_dir, date_order="dmy", debug=False):
    """One code folder's reports -> list of {"path", "date", "modality"}.

    The PDF name itself is preferred; when it has no date (or no modality)
    the enclosing folder names are searched instead, since the new file
    names don't always carry the code/date.
    """
    entries = []
    for path in sorted(code_dir.rglob("*.pdf")):
        if not path.is_file():
            continue
        date = modality = None
        date_from = modality_from = None
        for name in name_chain(path, code_dir):
            if date is None:
                date = parse_pdf_date(name, date_order)
                if date is not None:
                    date_from = name
            if modality is None:
                modality = parse_pdf_modality(name)
                if modality is not None:
                    modality_from = name
            if date is not None and modality is not None:
                break
        entries.append({"path": path, "date": date, "modality": modality})
        if debug:
            print(
                f"  {path.name!r} -> date={date!r} (from {date_from!r}) "
                f"modality={modality!r} (from {modality_from!r})"
            )
    return entries


def build_indices(code_reports):
    """entries -> (by_date_modality, by_date) lookup dicts for one code."""
    by_date_modality = defaultdict(list)
    by_date = defaultdict(list)
    for entry in code_reports:
        if entry["date"] is None:
            continue
        by_date[entry["date"]].append(entry)
        if entry["modality"] is not None:
            by_date_modality[(entry["date"], entry["modality"])].append(entry)
    return by_date_modality, by_date


def process_code(code, images_dirs, reports_dirs, date_order="dmy", debug=False):
    """Scan one code end to end, across every folder carrying that code.

    Returns (rows, stats).
    """
    started = time.time()
    images_dirs = list(images_dirs or [])
    reports_dirs = list(reports_dirs or [])

    code_reports = []
    for reports_dir in reports_dirs:
        code_reports.extend(
            scan_reports_dir(reports_dir, date_order=date_order, debug=debug)
        )
    by_date_modality, by_date = build_indices(code_reports)

    # (study_date, modality) -> {"count", "matched_pdfs"}
    summary = defaultdict(lambda: {"count": 0, "matched_pdfs": set()})
    matched_pdf_paths = set()
    files_seen = dicom_seen = 0

    for images_dir in images_dirs:
        for path in images_dir.rglob("*"):
            if not path.is_file():
                continue
            files_seen += 1
            if files_seen % PROGRESS_EVERY == 0:
                elapsed = time.time() - started
                rate = files_seen / elapsed if elapsed else 0
                log(
                    f"    {code}: {files_seen:,} files examined, "
                    f"{dicom_seen:,} DICOM, {format_duration(elapsed)} elapsed "
                    f"({rate:,.0f} files/s)"
                )
            info = read_dicom_study(path, debug=debug)
            if info is None:
                continue
            dicom_seen += 1
            study_date, modality = info

            if modality and (study_date, modality) in by_date_modality:
                matches = by_date_modality[(study_date, modality)]
            elif study_date in by_date:
                matches = by_date[study_date]
            else:
                matches = []

            bucket = summary[(study_date, modality)]
            bucket["count"] += 1
            for m in matches:
                matched_pdf_paths.add(m["path"])
                bucket["matched_pdfs"].add(str(m["path"]))

    rows = []
    matched = ambiguous = unmatched_images = 0
    for (study_date, modality), info in sorted(summary.items()):
        if info["matched_pdfs"]:
            status = "Matched"
            matched += 1
        elif code_reports:
            status = "Same folder, no date/modality match - review manually"
            ambiguous += 1
        else:
            status = "No PDF report found"
            unmatched_images += 1
        rows.append(
            [
                code,
                study_date,
                display_modality(modality),
                info["count"],
                ", ".join(sorted(info["matched_pdfs"])),
                status,
            ]
        )

    unmatched_pdf = 0
    for entry in sorted(
        code_reports, key=lambda e: (e["date"] or "", e["modality"] or "")
    ):
        if entry["path"] not in matched_pdf_paths:
            rows.append(
                [
                    code,
                    entry["date"] or "",
                    display_modality(entry["modality"]) or "",
                    "",
                    str(entry["path"]),
                    "No DICOM images found",
                ]
            )
            unmatched_pdf += 1

    stats = {
        "files_seen": files_seen,
        "dicom_seen": dicom_seen,
        "pdfs": len(code_reports),
        "matched": matched,
        "ambiguous": ambiguous,
        "unmatched_images": unmatched_images,
        "unmatched_pdf": unmatched_pdf,
        "elapsed": time.time() - started,
    }
    return rows, stats


def write_summary_sheet(path, rows, title="Summary"):
    """Write rows to a fresh workbook via a temp file, then swap it in.

    The swap matters: a half-written results file would otherwise look
    "done" to the next resumed run and get skipped.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(SUMMARY_HEADER)
    for row in rows:
        ws.append(row)
    for column, width in zip("ABCDEF", SUMMARY_WIDTHS):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    tmp = path.with_suffix(path.suffix + ".tmp")
    wb.save(tmp)
    os.replace(tmp, path)


def read_result_rows(path):
    """Read back one per-code results file, minus its header row."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        return [list(row) for row in ws.iter_rows(min_row=2, values_only=True)]
    finally:
        wb.close()


def combine_results(results_dir, output_xlsx):
    """Rebuild the all-codes summary from every per-code file on disk."""
    rows = []
    files = sorted(results_dir.glob("*-results.xlsx"))
    for path in files:
        try:
            rows.extend(read_result_rows(path))
        except Exception as exc:
            log(
                f"  WARNING: could not read {path.name} for the combined summary: {exc}"
            )
    write_summary_sheet(output_xlsx, rows)
    return len(files), len(rows)


def parse_args(argv):
    debug = "--debug" in argv
    force = "--force" in argv
    code_filter = None
    code_pattern = DEFAULT_CODE_PATTERN
    date_order = "dmy"
    workers = 1
    positional = []
    i = 0
    while i < len(argv):
        a = argv[i]
        if a in ("--debug", "--force"):
            i += 1
        elif a in ("--code", "--rcode"):
            if i + 1 >= len(argv):
                print(
                    "--code requires a value, e.g. --code AVSD0001 or "
                    "--code AVSD0001,AVSD0002"
                )
                sys.exit(1)
            code_filter = {
                c.strip().upper() for c in argv[i + 1].split(",") if c.strip()
            }
            i += 2
        elif a == "--code-regex":
            if i + 1 >= len(argv):
                print(r"--code-regex requires a value, e.g. --code-regex 'AVSD\d+'")
                sys.exit(1)
            code_pattern = argv[i + 1]
            i += 2
        elif a == "--date-order":
            if i + 1 >= len(argv):
                print("--date-order requires a value: dmy or mdy")
                sys.exit(1)
            date_order = argv[i + 1].strip().lower()
            if date_order not in ("dmy", "mdy"):
                print(f"--date-order must be dmy or mdy, got {argv[i + 1]!r}")
                sys.exit(1)
            i += 2
        elif a == "--workers":
            if i + 1 >= len(argv):
                print("--workers requires a value, e.g. --workers 4")
                sys.exit(1)
            try:
                workers = int(argv[i + 1])
            except ValueError:
                print(f"--workers value must be an integer, got {argv[i + 1]!r}")
                sys.exit(1)
            if workers < 1:
                print("--workers must be >= 1")
                sys.exit(1)
            i += 2
        else:
            positional.append(a)
            i += 1
    return positional, debug, force, code_filter, code_pattern, date_order, workers


def group_dirs_by_code(root, code_re, code_filter, side):
    """Subfolders of root -> {code: [dirs]}, ignoring extra text in the names.

    A folder whose name has no recognisable code keeps its full name as its
    key (and is reported), so nothing silently drops out of the run.
    """
    groups = defaultdict(list)
    nameless = []
    for path in sorted(p for p in root.iterdir() if p.is_dir()):
        code = extract_code(path.name, code_re)
        if code is None:
            code = path.name.strip().upper()
            nameless.append(path.name)
        if code_filter is not None and code not in code_filter:
            continue
        groups[code].append(path)
    if nameless:
        log(
            f"  WARNING: {len(nameless)} {side} folder(s) have no code in their "
            f"name and are keyed by folder name: {', '.join(nameless[:10])}"
            + (" ..." if len(nameless) > 10 else "")
        )
    return groups


def main():
    (
        args,
        debug,
        force,
        code_filter,
        code_pattern,
        date_order,
        workers,
    ) = parse_args(sys.argv[1:])

    if len(args) < 2:
        print(__doc__)
        sys.exit(1)

    try:
        code_re = re.compile(code_pattern, re.IGNORECASE)
    except re.error as exc:
        print(f"--code-regex is not a valid regular expression: {exc}")
        sys.exit(1)

    images_root = Path(args[0]).expanduser().resolve()
    reports_root = Path(args[1]).expanduser().resolve()
    if not images_root.is_dir():
        print(f"Not a folder: {images_root}")
        sys.exit(1)
    if not reports_root.is_dir():
        print(f"Not a folder: {reports_root}")
        sys.exit(1)

    output_xlsx = Path(args[2]) if len(args) > 2 else Path.cwd() / "match_report.xlsx"
    results_dir = output_xlsx.parent / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    open_log(results_dir / "progress.log")

    log(f"Images:  {images_root}")
    log(f"Reports: {reports_root}")
    log(f"Results: {results_dir}")
    log(f"Code pattern: {code_pattern}  |  ambiguous numeric dates read as {date_order}")

    image_dirs = group_dirs_by_code(images_root, code_re, code_filter, "images")
    report_dirs = group_dirs_by_code(reports_root, code_re, code_filter, "reports")
    codes = sorted(set(image_dirs) | set(report_dirs))

    if code_filter is not None:
        log(f"Restricting scan to code(s): {sorted(code_filter)}")

    if debug:
        print("\n--- code comparison ---")
        print(f"In both: {sorted(set(image_dirs) & set(report_dirs))}")
        print(f"Only in images: {sorted(set(image_dirs) - set(report_dirs))}")
        print(f"Only in reports: {sorted(set(report_dirs) - set(image_dirs))}")
        for name in codes:
            print(f"  code: {name!r}")
            for path in image_dirs.get(name, []):
                print(f"      images  <- {path.name!r}")
            for path in report_dirs.get(name, []):
                print(f"      reports <- {path.name!r}")
        print("--- end code comparison ---\n")

    todo = [c for c in codes if force or not result_path(results_dir, c).exists()]
    already_done = len(codes) - len(todo)
    log(
        f"{len(codes)} code(s) in scope; {already_done} already have results, {len(todo)} to do"
    )
    if already_done and not force:
        log("(pass --force to re-scan the ones that already have a results file)")

    run_started = time.time()
    totals = defaultdict(int)

    if workers == 1:
        for index, code in enumerate(todo, start=1):
            remaining = len(todo) - index
            log(
                f"[{index}/{len(todo)}] {code}: starting ({remaining} code(s) left after this)"
            )
            rows, stats = process_code(
                code,
                image_dirs.get(code),
                report_dirs.get(code),
                date_order=date_order,
                debug=debug,
            )
            out = result_path(results_dir, code)
            write_summary_sheet(out, rows)
            for key, value in stats.items():
                totals[key] += value
            log(
                f"[{index}/{len(todo)}] {code}: done in {format_duration(stats['elapsed'])} - "
                f"{stats['files_seen']:,} files ({stats['dicom_seen']:,} DICOM), "
                f"{stats['pdfs']} PDF(s), {stats['matched']} matched study bucket(s), "
                f"{stats['ambiguous']} needing review, {stats['unmatched_pdf']} unmatched PDF(s) "
                f"-> {out.name}"
            )
            elapsed = time.time() - run_started
            if remaining:
                eta = elapsed / index * remaining
                log(
                    f"    run elapsed {format_duration(elapsed)}, rough ETA {format_duration(eta)} for the rest"
                )
    else:
        log(f"Running with {workers} parallel workers, one code per worker")
        completed = 0
        with concurrent.futures.ProcessPoolExecutor(
            max_workers=workers,
            initializer=_init_worker,
            initargs=(results_dir / "progress.log",),
        ) as pool:
            futures = {
                pool.submit(
                    process_code,
                    code,
                    image_dirs.get(code),
                    report_dirs.get(code),
                    date_order,
                    debug,
                ): code
                for code in todo
            }
            for future in concurrent.futures.as_completed(futures):
                code = futures[future]
                completed += 1
                remaining = len(todo) - completed
                try:
                    rows, stats = future.result()
                except Exception as exc:
                    log(f"[{completed}/{len(todo)}] {code}: FAILED: {type(exc).__name__}: {exc}")
                    continue
                out = result_path(results_dir, code)
                write_summary_sheet(out, rows)
                for key, value in stats.items():
                    totals[key] += value
                log(
                    f"[{completed}/{len(todo)}] {code}: done in {format_duration(stats['elapsed'])} - "
                    f"{stats['files_seen']:,} files ({stats['dicom_seen']:,} DICOM), "
                    f"{stats['pdfs']} PDF(s), {stats['matched']} matched study bucket(s), "
                    f"{stats['ambiguous']} needing review, {stats['unmatched_pdf']} unmatched PDF(s) "
                    f"-> {out.name}"
                )
                elapsed = time.time() - run_started
                if remaining:
                    eta = elapsed / completed * remaining
                    log(
                        f"    run elapsed {format_duration(elapsed)}, rough ETA {format_duration(eta)} for the rest"
                    )

    log("Rebuilding combined summary from results/ ...")
    file_count, row_count = combine_results(results_dir, output_xlsx)

    log("")
    log(f"Codes scanned this run: {len(todo)} (skipped as already done: {already_done})")
    log(
        f"Files examined: {totals['files_seen']:,} ({totals['dicom_seen']:,} readable DICOM)"
    )
    log(f"Matched study buckets: {totals['matched']}")
    log(f"Same folder, no date/modality match (needs review): {totals['ambiguous']}")
    log(f"DICOM studies with no PDF report: {totals['unmatched_images']}")
    log(f"PDF reports with no DICOM images: {totals['unmatched_pdf']}")
    log(f"Total run time: {format_duration(time.time() - run_started)}")
    log(f"Per-code results: {results_dir}")
    log(
        f"Combined summary ({row_count:,} rows from {file_count} code file(s)): {output_xlsx}"
    )


if __name__ == "__main__":
    main()
