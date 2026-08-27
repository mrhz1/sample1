"""Merge every per-code results workbook into one master Excel file.

Reads all `<code>-results.xlsx` files produced by match_reports.py and writes
a single workbook with two sheets:

  "All Data" - every row from every file, same columns and same format.
  "Totals"   - plain counts: matched vs not matched, and a per-modality
               breakdown.

Usage:
    python merge_results.py [results_dir] [output.xlsx]

    results_dir    Folder holding the *-results.xlsx files (default: ./results)
    output.xlsx    Master workbook to write (default: <results_dir>/../master_report.xlsx)

Examples:
    python merge_results.py
    python merge_results.py results master_report.xlsx
    python merge_results.py "/mnt/data/results" "/mnt/data/master_report.xlsx"
"""

import os
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

# Must stay in step with match_reports.py's SUMMARY_HEADER / SUMMARY_WIDTHS.
SUMMARY_HEADER = [
    "Code",
    "Study Date",
    "Modality",
    "DICOM File Count",
    "Matched PDF File(s)",
    "Match Status",
]
SUMMARY_WIDTHS = (12, 14, 10, 16, 45, 40)

MATCHED_STATUS = "Matched"
NO_DICOM_STATUS = "No DICOM images found"

BOLD = Font(bold=True)


def read_result_rows(path):
    """Read back one per-code results file, minus its header row."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        return [list(row) for row in ws.iter_rows(min_row=2, values_only=True)]
    finally:
        wb.close()


def to_int(value):
    """DICOM File Count is blank on PDF-only rows; treat that as 0."""
    if value in (None, ""):
        return 0
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def collect_rows(results_dir):
    """Load every per-code workbook in the folder, sorted by filename."""
    rows = []
    files = sorted(results_dir.glob("*-results.xlsx"))
    if not files:
        print(f"No *-results.xlsx files found in {results_dir}")
        return rows, 0
    for path in files:
        try:
            file_rows = read_result_rows(path)
        except Exception as exc:
            print(f"  WARNING: could not read {path.name}: {exc}")
            continue
        print(f"  {path.name}: {len(file_rows):,} row(s)")
        rows.extend(file_rows)
    return rows, len(files)


def tally(rows):
    """Boil the merged rows down to the numbers for the Totals sheet."""
    totals = {
        "rows": len(rows),
        "codes": len({r[0] for r in rows if r[0]}),
        "dicom_total": 0,
        "dicom_matched": 0,
        "dicom_unmatched": 0,
        "studies_matched": 0,
        "studies_unmatched": 0,
        "pdf_only_rows": 0,
    }
    statuses = Counter()
    per_modality = defaultdict(
        lambda: {
            "studies": 0,
            "dicom": 0,
            "dicom_matched": 0,
            "dicom_unmatched": 0,
            "pdf_only": 0,
        }
    )

    for row in rows:
        modality = (row[2] or "(blank)").strip() or "(blank)"
        count = to_int(row[3])
        status = (row[5] or "").strip()
        statuses[status or "(blank)"] += 1

        bucket = per_modality[modality]
        bucket["dicom"] += count
        totals["dicom_total"] += count

        if status == NO_DICOM_STATUS:
            totals["pdf_only_rows"] += 1
            bucket["pdf_only"] += 1
            continue

        bucket["studies"] += 1
        if status == MATCHED_STATUS:
            totals["studies_matched"] += 1
            totals["dicom_matched"] += count
            bucket["dicom_matched"] += count
        else:
            totals["studies_unmatched"] += 1
            totals["dicom_unmatched"] += count
            bucket["dicom_unmatched"] += count

    return totals, statuses, per_modality


def write_master(path, rows, file_count):
    """Write the master workbook via a temp file, then swap it in."""
    totals, statuses, per_modality = tally(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "All Data"
    ws.append(SUMMARY_HEADER)
    for cell in ws[1]:
        cell.font = BOLD
    for row in rows:
        ws.append(row)
    for column, width in zip("ABCDEF", SUMMARY_WIDTHS):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"

    ts = wb.create_sheet("Totals")
    ts.column_dimensions["A"].width = 42
    ts.column_dimensions["B"].width = 14
    ts.column_dimensions["C"].width = 14
    ts.column_dimensions["D"].width = 16

    def section(title):
        ts.append([])
        ts.append([title])
        ts.cell(row=ts.max_row, column=1).font = BOLD

    ts.append(["Overall"])
    ts.cell(row=1, column=1).font = BOLD
    ts.append(["Result files merged", file_count])
    ts.append(["Codes", totals["codes"]])
    ts.append(["Rows", totals["rows"]])
    ts.append(["DICOM files total", totals["dicom_total"]])
    ts.append(["DICOM files matched", totals["dicom_matched"]])
    ts.append(["DICOM files not matched", totals["dicom_unmatched"]])
    ts.append(["Studies matched", totals["studies_matched"]])
    ts.append(["Studies not matched", totals["studies_unmatched"]])
    ts.append(["PDF reports with no DICOM", totals["pdf_only_rows"]])

    section("By match status")
    ts.append(["Status", "Rows"])
    for cell in ts[ts.max_row]:
        cell.font = BOLD
    for status, count in statuses.most_common():
        ts.append([status, count])

    section("By modality")
    ts.append(
        [
            "Modality",
            "Studies",
            "DICOM files",
            "DICOM matched",
            "DICOM not matched",
            "PDF only (no DICOM)",
        ]
    )
    for cell in ts[ts.max_row]:
        cell.font = BOLD
    ts.column_dimensions["E"].width = 18
    ts.column_dimensions["F"].width = 20
    for modality in sorted(per_modality):
        b = per_modality[modality]
        ts.append(
            [
                modality,
                b["studies"],
                b["dicom"],
                b["dicom_matched"],
                b["dicom_unmatched"],
                b["pdf_only"],
            ]
        )

    tmp = path.with_suffix(path.suffix + ".tmp")
    wb.save(tmp)
    os.replace(tmp, path)
    return totals


def main(argv):
    if argv and argv[0] in ("-h", "--help"):
        print(__doc__)
        return 0

    results_dir = Path(argv[0]) if argv else Path.cwd() / "results"
    if not results_dir.is_dir():
        print(f"Results folder not found: {results_dir}")
        return 1

    if len(argv) > 1:
        output_xlsx = Path(argv[1])
    else:
        output_xlsx = results_dir.parent / "master_report.xlsx"
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)

    print(f"Reading: {results_dir}")
    rows, file_count = collect_rows(results_dir)
    if not rows:
        print("Nothing to merge.")
        return 1

    totals = write_master(output_xlsx, rows, file_count)

    print("")
    print(f"Master file: {output_xlsx}")
    print(f"  files merged            {file_count:,}")
    print(f"  codes                   {totals['codes']:,}")
    print(f"  rows                    {totals['rows']:,}")
    print(f"  DICOM files total       {totals['dicom_total']:,}")
    print(f"  DICOM files matched     {totals['dicom_matched']:,}")
    print(f"  DICOM files not matched {totals['dicom_unmatched']:,}")
    print(f"  studies matched         {totals['studies_matched']:,}")
    print(f"  studies not matched     {totals['studies_unmatched']:,}")
    print(f"  PDFs with no DICOM      {totals['pdf_only_rows']:,}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
